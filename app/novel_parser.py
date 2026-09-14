"""
小说解析器 —— 支持任意常见格式小说文件的上传解析

覆盖格式：
- txt / text / log：自动探测编码（utf-8 / utf-8-sig / GBK / GB18030 / BIG5 / utf-16），
  中文小说常见的 GBK 系列编码优先使用 charset-normalizer 统计判定，并用「严格 utf-8 解码」作为第一判据
- md / markdown：按纯文本处理（章节支持 `# 标题`）
- docx：python-docx 抽取段落 + 表格
- pdf：pypdf 逐页抽取（扫描版 PDF 会给出 OCR 提示）
- epub：ebooklib 抽取正文（缺失时回退 zipfile 解析）
- html / htm / xhtml：标准库 HTMLParser 去标签抽取正文

统一输出：纯文本 + 字符统计 + 章节切分，并落盘到项目 novels 目录。
"""
from __future__ import annotations

import json
import logging
import os
import re
import time
import uuid
from datetime import datetime
from html import unescape as html_unescape
from html.parser import HTMLParser

logger = logging.getLogger(__name__)

# 支持的扩展名（.doc 为旧版二进制格式，单独给出友好错误提示）
TEXT_EXTS = {".txt", ".text", ".log", ".md", ".markdown"}
DOC_EXTS = {".docx", ".pdf", ".epub", ".html", ".htm", ".xhtml"}
SUPPORTED_EXTS = TEXT_EXTS | DOC_EXTS
LEGACY_EXTS = {".doc", ".wps", ".rtf", ".mobi", ".azw", ".azw3"}

MAX_CHAPTERS_KEPT = 3000


class NovelParseError(Exception):
    """小说解析失败（面向用户的友好错误）"""


# ===================== 编码探测 =====================

def _han_ratio(text: str) -> float:
    if not text:
        return 0.0
    sample = text[:200000]
    han = sum(1 for ch in sample if "\u4e00" <= ch <= "\u9fff")
    return han / max(len(sample), 1)


def _ascii_ratio(text: str) -> float:
    if not text:
        return 0.0
    sample = text[:200000]
    asc = sum(1 for ch in sample if ord(ch) < 128)
    return asc / max(len(sample), 1)


def _bad_ratio(text: str) -> float:
    if not text:
        return 1.0
    sample = text[:200000]
    bad = sample.count("\ufffd") + sum(
        1 for ch in sample if ord(ch) < 32 and ch not in "\n\r\t"
    )
    return bad / max(len(sample), 1)


def decode_bytes(blob: bytes):
    """把字节流解码为文本，返回 (text, encoding, note)。

    策略（中文小说友好）：
    1. BOM 判定（utf-8-sig / utf-16）
    2. 严格 utf-8 解码成功 → utf-8（GBK 字节几乎不可能整体通过 utf-8 严格校验）
    3. charset-normalizer 统计判定（能正确识别 GBK / GB18030 / BIG5 等）
    4. 依次回退 gb18030 → big5 → cp936，用「汉字占比 / ASCII 占比 / 坏字符占比」打分择优
    """
    if blob.startswith(b"\xef\xbb\xbf"):
        return blob.decode("utf-8-sig", errors="replace"), "utf-8-sig", "BOM 识别为 UTF-8(BOM)"
    if blob[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return blob.decode("utf-16", errors="replace"), "utf-16", "BOM 识别为 UTF-16"

    try:
        text = blob.decode("utf-8")
        if "\ufffd" not in text and _bad_ratio(text) < 0.001:
            return text, "utf-8", "严格 UTF-8 解码成功"
    except UnicodeDecodeError:
        pass

    try:
        from charset_normalizer import from_bytes

        best = from_bytes(blob).best()
        if best is not None and best.encoding:
            enc = str(best.encoding).replace("_", "-")
            try:
                text = blob.decode(enc, errors="replace")
                if _bad_ratio(text) < 0.005:
                    return text, enc, f"charset-normalizer 统计判定为 {enc}"
            except (LookupError, UnicodeDecodeError):
                pass
    except ImportError:
        logger.warning("charset-normalizer 未安装，中文编码将使用回退策略")
    except Exception as e:  # noqa: BLE001
        logger.warning(f"charset-normalizer 探测异常：{e}")

    best_pair = (None, "", -1.0)
    for enc in ("gb18030", "big5", "cp936"):
        try:
            text = blob.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
        score = _han_ratio(text) * 2 + _ascii_ratio(text) - _bad_ratio(text) * 10
        if score > best_pair[2]:
            best_pair = (text, enc, score)
    if best_pair[0] is not None:
        return best_pair[0], best_pair[1], f"回退解码（{best_pair[1]}），适用于 GBK/GB2312/GB18030 中文小说"

    text = blob.decode("gb18030", errors="replace")
    return text, "gb18030(errors=replace)", "强制 GB18030 解码，部分字符可能丢失"


# ===================== 各格式抽取 =====================

class _HTMLTextExtractor(HTMLParser):
    """去掉 script/style，抽取正文文本"""

    # 注意：只放"成对出现"的容器标签；meta / link / br 等空元素无闭合标签，
    # 若放进 _SKIP 会让跳过深度只增不减，从而把整篇正文全部丢弃。
    _SKIP = {"script", "style", "noscript", "head", "svg", "iframe", "template", "canvas", "form"}
    _BLOCK = {
        "p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6",
        "section", "article", "blockquote", "pre", "td", "th", "hr",
    }

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        tag = (tag or "").lower()
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        tag = (tag or "").lower()
        if tag in self._SKIP and self._skip_depth > 0:
            self._skip_depth -= 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._skip_depth == 0 and data:
            self.parts.append(data)

    def text(self) -> str:
        return "".join(self.parts)


def html_to_text(html: str) -> str:
    parser = _HTMLTextExtractor()
    text = ""
    try:
        parser.feed(html)
        text = parser.text()
    except Exception as e:  # noqa: BLE001
        logger.warning(f"HTMLParser 解析异常，改用正则清洗：{e}")
    if len(re.sub(r"\s", "", text)) < 20:
        # 兜底：异常/非常规 HTML（如残缺标签、大写标签、编码声明错乱）改用正则清洗
        cleaned = re.sub(r"(?is)<(script|style|noscript|head|svg).*?</\1>", " ", html)
        cleaned = re.sub(r"(?s)<[^>]+>", "\n", cleaned)
        cleaned = html_unescape(cleaned)
        if len(re.sub(r"\s", "", cleaned)) > len(re.sub(r"\s", "", text)):
            text = cleaned
    return text


def _read_text_file(path: str):
    with open(path, "rb") as f:
        blob = f.read()
    text, enc, note = decode_bytes(blob)
    return text, enc, note, []


def _read_docx(path: str):
    try:
        import docx  # python-docx
    except ImportError as e:  # pragma: no cover
        raise NovelParseError("缺少 python-docx 依赖，无法解析 .docx") from e
    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts), "docx(xml)", "python-docx 抽取段落与表格", []


def _read_pdf(path: str):
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise NovelParseError("缺少 pypdf 依赖，无法解析 .pdf") from e
    reader = PdfReader(path)
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as e:  # noqa: BLE001
            logger.warning(f"PDF 某页抽取失败：{e}")
            pages.append("")
    text = "\n\n".join(pages)
    warnings = []
    if len(text.strip()) < 50:
        warnings.append("PDF 抽取到的文本极少，可能是扫描版图片 PDF，需先做 OCR 才能转剧本")
    return text, "pdf", f"pypdf 抽取 {len(reader.pages)} 页", warnings


def _read_epub(path: str):
    warnings = []
    try:
        import ebooklib
        from ebooklib import epub

        book = epub.read_epub(path)
        parts = []
        for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
            try:
                raw = item.get_content().decode("utf-8", errors="replace")
            except Exception:  # noqa: BLE001
                continue
            plain = html_to_text(raw).strip()
            if plain:
                parts.append(plain)
        if parts:
            return "\n\n".join(parts), "epub", f"ebooklib 抽取 {len(parts)} 个文档节点", warnings
        warnings.append("ebooklib 未抽到正文，已回退 zipfile 解析")
    except ImportError:
        warnings.append("未安装 ebooklib，已回退 zipfile 解析")
    except Exception as e:  # noqa: BLE001
        warnings.append(f"ebooklib 解析异常（{e}），已回退 zipfile 解析")

    import zipfile

    with zipfile.ZipFile(path) as zf:
        names = [
            n for n in zf.namelist()
            if n.lower().endswith((".xhtml", ".html", ".htm")) and "META-INF" not in n
        ]
        names.sort()
        parts = []
        for n in names:
            raw = zf.read(n).decode("utf-8", errors="replace")
            plain = html_to_text(raw).strip()
            if plain:
                parts.append(plain)
    if not parts:
        raise NovelParseError("EPUB 中未找到可读正文")
    return "\n\n".join(parts), "epub(zip)", f"zipfile 抽取 {len(parts)} 个文档节点", warnings


def _read_html(path: str):
    with open(path, "rb") as f:
        blob = f.read()
    html, enc, note = decode_bytes(blob)
    text = html_to_text(html)
    return text, enc, f"HTML 去标签（{note}）", []


def normalize_text(text: str) -> str:
    text = (text or "").replace("\ufeff", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u3000]+$", "", text, flags=re.M)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    return text.strip()


# ===================== 章节切分 =====================

_CN_NUM = "0-9０-９零一二三四五六七八九十百千万两〇"
CHAPTER_PATTERNS = [
    re.compile(rf"^[ \t\u3000]*第[ \t]*[{_CN_NUM}]{{1,12}}[ \t]*[章回节卷篇集部][ \t]*[：:.、·\-—－]?[ \t]*(.{{0,50}})$", re.M),
    re.compile(r"^[ \t\u3000]*(?:Chapter|CHAPTER|chapter)[ \t]+\d{1,5}[ \t]*[：:.、\-]?[ \t]*(.{0,50})$", re.M),
    re.compile(r"^[ \t\u3000]*#{1,4}[ \t]+\S.{0,60}$", re.M),
    re.compile(r"^[ \t\u3000]*(?:序章|楔子|前言|引子|后记|尾声|终章|大结局|番外).{0,30}$", re.M),
]


def split_chapters(text: str):
    """切分章节，返回 [{'index','title','start','end','char_count'}]（不含正文，避免重复占用内存）"""
    if not text:
        return []
    marks = []
    for pat in CHAPTER_PATTERNS:
        for m in pat.finditer(text):
            line = m.group(0).strip()
            if not line or len(line) > 70:
                continue
            marks.append((m.start(), line[:70]))
    if not marks:
        return []

    marks.sort()
    dedup = []
    seen_lines = set()
    for pos, title in marks:
        line_no = text.count("\n", 0, pos)
        if line_no in seen_lines:
            continue
        seen_lines.add(line_no)
        title = re.sub(r"^#+[ \t]*", "", title).strip()
        dedup.append((pos, title or f"第{len(dedup) + 1}节"))
        if len(dedup) >= MAX_CHAPTERS_KEPT:
            break

    chapters = []
    for i, (pos, title) in enumerate(dedup):
        end = dedup[i + 1][0] if i + 1 < len(dedup) else len(text)
        chapters.append({
            "index": i + 1,
            "title": title,
            "start": pos,
            "end": end,
            "char_count": end - pos,
        })
    return chapters


def guess_title(text: str, fallback: str) -> str:
    for raw in text.split("\n")[:60]:
        line = raw.strip().strip("#").strip()
        if not line:
            continue
        if any(p.match(line) for p in CHAPTER_PATTERNS):
            continue
        if 2 <= len(line) <= 30:
            return line
    return fallback


# ===================== 对外主流程 =====================

def extract_text(path: str, display_name: str = None) -> dict:
    """从任意支持格式抽取纯文本，返回 {text, ext, format, encoding, encoding_note, warnings}"""
    name = display_name or os.path.basename(path)
    ext = os.path.splitext(name)[1].lower()
    if ext in LEGACY_EXTS:
        raise NovelParseError(
            f"暂不支持 {ext} 格式（旧版二进制/私有格式），请先另存为 .docx / .txt / .pdf 后再上传"
        )
    if ext in TEXT_EXTS:
        text, enc, note, warns = _read_text_file(path)
        fmt = "markdown" if ext in (".md", ".markdown") else "text"
    elif ext == ".docx":
        text, enc, note, warns = _read_docx(path)
        fmt = "docx"
    elif ext == ".pdf":
        text, enc, note, warns = _read_pdf(path)
        fmt = "pdf"
    elif ext == ".epub":
        text, enc, note, warns = _read_epub(path)
        fmt = "epub"
    elif ext in (".html", ".htm", ".xhtml"):
        text, enc, note, warns = _read_html(path)
        fmt = "html"
    else:
        # 未知扩展名：按纯文本尽力解析
        text, enc, note, warns = _read_text_file(path)
        fmt = ext.lstrip(".") or "text"
        warns = list(warns) + [f"未知扩展名 {ext}，已按纯文本尽力解析"]

    return {
        "text": text,
        "ext": ext,
        "format": fmt,
        "encoding": enc,
        "encoding_note": note,
        "warnings": warns,
    }


def ingest_novel(raw_path: str, filename: str, novels_dir: str) -> dict:
    """解析上传的小说原始文件 → 落盘标准化文本 + 元数据，返回元数据 dict"""
    os.makedirs(novels_dir, exist_ok=True)
    info = extract_text(raw_path, filename)
    text = normalize_text(info["text"])
    if len(text.strip()) < 20:
        raise NovelParseError("未能从文件中提取到有效文本（内容过短或为扫描件）")

    chapters = split_chapters(text)
    novel_id = time.strftime("%Y%m%d_%H%M%S") + "_" + uuid.uuid4().hex[:6]
    text_name = f"{novel_id}.txt"
    text_path = os.path.join(novels_dir, text_name)
    with open(text_path, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)

    stem = os.path.splitext(os.path.basename(filename))[0]
    meta = {
        "novel_id": novel_id,
        "name": stem or novel_id,
        "title": guess_title(text, stem or novel_id),
        "source_filename": os.path.basename(filename),
        "ext": info["ext"],
        "format": info["format"],
        "encoding": info["encoding"],
        "encoding_note": info["encoding_note"],
        "warnings": info["warnings"],
        "size_bytes": os.path.getsize(raw_path),
        "total_chars": len(text),
        "char_count": len(re.sub(r"\s", "", text)),
        "line_count": text.count("\n") + 1,
        "chapter_count": len(chapters),
        "chapters": chapters[:MAX_CHAPTERS_KEPT],
        "text_file": text_name,
        "uploaded_at": datetime.now().isoformat(timespec="seconds"),
    }
    with open(os.path.join(novels_dir, f"{novel_id}.json"), "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)

    index = _load_index(novels_dir)
    index = [m for m in index if m.get("novel_id") != novel_id]
    index.append({k: v for k, v in meta.items() if k != "chapters"})
    index.sort(key=lambda m: m.get("uploaded_at", ""), reverse=True)
    _save_index(novels_dir, index)
    logger.info(f"小说已入库：{novel_id}（{meta['char_count']} 字 / {meta['chapter_count']} 章）")
    return meta


# ===================== 索引与读取 =====================

def _index_path(novels_dir: str) -> str:
    return os.path.join(novels_dir, "index.json")


def _load_index(novels_dir: str) -> list:
    path = _index_path(novels_dir)
    if not os.path.isfile(path):
        return []
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, list) else []
    except Exception as e:  # noqa: BLE001
        logger.warning(f"小说索引读取失败：{e}")
        return []


def _save_index(novels_dir: str, index: list):
    os.makedirs(novels_dir, exist_ok=True)
    with open(_index_path(novels_dir), "w", encoding="utf-8") as f:
        json.dump(index, f, ensure_ascii=False, indent=2)


def list_novels(novels_dir: str) -> list:
    index = _load_index(novels_dir)
    # 过滤掉实体文件已丢失的条目
    alive = []
    for m in index:
        text_file = m.get("text_file") or f"{m.get('novel_id')}.txt"
        if os.path.isfile(os.path.join(novels_dir, text_file)):
            alive.append(m)
    if len(alive) != len(index):
        _save_index(novels_dir, alive)
    return alive


def get_novel(novels_dir: str, novel_id: str) -> dict:
    path = os.path.join(novels_dir, f"{novel_id}.json")
    if not os.path.isfile(path):
        raise NovelParseError(f"小说不存在：{novel_id}")
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def read_novel_text(novels_dir: str, novel_id: str, max_chars: int = None) -> str:
    meta = get_novel(novels_dir, novel_id)
    text_path = os.path.join(novels_dir, meta.get("text_file") or f"{novel_id}.txt")
    if not os.path.isfile(text_path):
        raise NovelParseError(f"小说正文文件缺失：{text_path}")
    with open(text_path, "r", encoding="utf-8") as f:
        text = f.read()
    if max_chars and max_chars > 0:
        return text[:max_chars]
    return text


def preview_novel(novels_dir: str, novel_id: str, offset: int = 0, limit: int = 4000) -> dict:
    meta = get_novel(novels_dir, novel_id)
    text = read_novel_text(novels_dir, novel_id)
    offset = max(0, int(offset or 0))
    limit = max(200, min(int(limit or 4000), 20000))
    chunk = text[offset:offset + limit]
    return {
        "novel_id": novel_id,
        "name": meta.get("name"),
        "title": meta.get("title"),
        "offset": offset,
        "limit": limit,
        "text": chunk,
        "total_chars": len(text),
        "has_more": offset + limit < len(text),
    }
